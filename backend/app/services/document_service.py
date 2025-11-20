"""
Document processing service for knowledge base
"""
import os
from typing import List, BinaryIO
from pathlib import Path
import PyPDF2
import docx
import google.generativeai as genai
from sqlalchemy.ext.asyncio import AsyncSession
from app.models.document import Document, KnowledgeChunk
from app.models.user import User
from app.core.config import settings


class DocumentService:
    """Service for document processing and embedding"""
    
    def __init__(self, db: AsyncSession):
        self.db = db
        
        # Configure Gemini
        genai.configure(api_key=settings.GEMINI_API_KEY)
        self.embedding_model = 'models/embedding-001'
        
        self.upload_dir = Path(settings.UPLOAD_DIR)
        self.upload_dir.mkdir(exist_ok=True)
    
    def extract_text_from_pdf(self, file_path: Path) -> str:
        """Extract text from PDF file"""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        except Exception as e:
            raise ValueError(f"Failed to extract text from PDF: {str(e)}")
        return text
    
    def extract_text_from_docx(self, file_path: Path) -> str:
        """Extract text from DOCX file"""
        try:
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
        return text
    
    def extract_text_from_txt(self, file_path: Path) -> str:
        """Extract text from TXT file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
        except Exception as e:
            raise ValueError(f"Failed to read text file: {str(e)}")
        return text
    
    def chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
        """Split text into overlapping chunks"""
        words = text.split()
        chunks = []
        
        for i in range(0, len(words), chunk_size - overlap):
            chunk = ' '.join(words[i:i + chunk_size])
            if chunk.strip():
                chunks.append(chunk)
        
        return chunks if chunks else [text]  # Return original text if no chunks
    
    async def generate_embedding(self, text: str) -> List[float]:
        """Generate embedding for text"""
        try:
            result = genai.embed_content(
                model=self.embedding_model,
                content=text[:10000],  # Limit text length
                task_type="retrieval_document"
            )
            return result['embedding']
        except Exception as e:
            raise ValueError(f"Failed to generate embedding: {str(e)}")
    
    async def process_document(
        self,
        file_content: BinaryIO,
        filename: str,
        title: str,
        user: User,
        description: str = None,
        department: str = None
    ) -> Document:
        """
        Process uploaded document and create embeddings
        
        Args:
            file_content: File content
            filename: Original filename
            title: Document title
            user: Uploader user
            description: Optional description
            department: Optional department
            
        Returns:
            Created document
        """
        # Validate file type
        file_ext = Path(filename).suffix.lower()
        if file_ext not in settings.allowed_file_extensions:
            raise ValueError(f"File type {file_ext} not allowed")
        
        # Save file
        safe_filename = f"{user.id}_{filename}"
        file_path = self.upload_dir / safe_filename
        
        with open(file_path, 'wb') as f:
            f.write(file_content.read())
        
        file_size = file_path.stat().st_size
        
        # Validate file size
        if file_size > settings.max_file_size_bytes:
            file_path.unlink()  # Delete file
            raise ValueError(f"File size exceeds maximum of {settings.MAX_FILE_SIZE_MB} MB")
        
        # Create document record
        document = Document(
            title=title,
            description=description,
            file_path=str(file_path),
            file_name=filename,
            file_size=file_size,
            file_type=file_ext,
            department=department,
            uploaded_by_id=user.id,
            status="processing"
        )
        
        self.db.add(document)
        await self.db.commit()
        await self.db.refresh(document)
        
        try:
            # Extract text based on file type
            if file_ext == '.pdf':
                text = self.extract_text_from_pdf(file_path)
            elif file_ext == '.docx':
                text = self.extract_text_from_docx(file_path)
            elif file_ext == '.txt':
                text = self.extract_text_from_txt(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_ext}")
            
            # Chunk text
            chunks = self.chunk_text(text)
            
            # Create embeddings for each chunk
            for i, chunk_text in enumerate(chunks):
                embedding = await self.generate_embedding(chunk_text)
                
                chunk = KnowledgeChunk(
                    document_id=document.id,
                    chunk_index=i,
                    content=chunk_text,
                    embedding=embedding,
                    chunk_metadata={'chunk_size': len(chunk_text)}
                )
                
                self.db.add(chunk)
            
            # Update document status
            document.status = "ready"
            document.chunks_count = len(chunks)
            document.embeddings_count = len(chunks)
            
            await self.db.commit()
            await self.db.refresh(document)
            
            return document
        except Exception as e:
            # Update document status to failed
            document.status = "failed"
            await self.db.commit()
            raise ValueError(f"Failed to process document: {str(e)}")

    async def process_document_content(
        self,
        document_id,
        content: str,
        file_type: str = 'web'
    ) -> None:
        """
        Process text content and create embeddings

        Args:
            document_id: Document ID
            content: Text content to process
            file_type: Type of content (default: 'web')
        """
        try:
            # Chunk text
            chunks = self.chunk_text(content)

            # Create embeddings for each chunk
            for i, chunk_text in enumerate(chunks):
                embedding = await self.generate_embedding(chunk_text)

                chunk = KnowledgeChunk(
                    document_id=document_id,
                    chunk_index=i,
                    content=chunk_text,
                    embedding=embedding,
                    chunk_metadata={'chunk_size': len(chunk_text), 'source_type': file_type}
                )

                self.db.add(chunk)

            await self.db.commit()

        except Exception as e:
            await self.db.rollback()
            raise ValueError(f"Failed to process content: {str(e)}")
